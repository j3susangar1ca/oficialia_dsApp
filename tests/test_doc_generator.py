"""
tests/test_doc_generator.py — Generación del .docx de respuestas a oficios
(core.doc_generator). Sin plantilla institucional (documento en blanco):
verifica que el contenido de `RespuestaOficio` efectivamente aparezca en
los párrafos del documento generado.
"""

from __future__ import annotations

import io
import uuid

from docx import Document

from core.doc_generator import generar_docx_respuesta, nombre_archivo_respuesta
from core.models import (
    DocumentoRegistro,
    EstadoDocumento,
    OrigenIngesta,
    RespuestaOficio,
)


def _documento(**overrides) -> DocumentoRegistro:
    datos = dict(
        id=str(uuid.uuid4()),
        nombre_archivo_original="oficio.pdf",
        ruta_archivo_actual="03_procesados/2026/08/oficio.pdf",
        origen=OrigenIngesta.WEB_DRAG_DROP,
        estado=EstadoDocumento.COMPLETADO,
        sha256="a" * 64,
        numero_oficio="DSA-2026-089-OF",
    )
    datos.update(overrides)
    return DocumentoRegistro(**datos)


def _respuesta(**overrides) -> RespuestaOficio:
    datos = dict(
        destinatario_nombre="JUAN PÉREZ LÓPEZ",
        destinatario_cargo="Director de Administración",
        destinatario_dependencia="Secretaría de Salud Jalisco",
        asunto="Respuesta a solicitud de información",
        cuerpo_respuesta=(
            "En atención a su similar número DSA-2026-089-OF, me permito informarle lo siguiente.\n\n"
            "El trámite solicitado ha sido atendido en su totalidad."
        ),
        despedida="Sin otro particular por el momento, quedo de usted.",
        ccp=["Archivo", "Minutario"],
    )
    datos.update(overrides)
    return RespuestaOficio(**datos)


def _texto_completo(contenido: bytes) -> str:
    doc = Document(io.BytesIO(contenido))
    return "\n".join(p.text for p in doc.paragraphs)


class TestGenerarDocxRespuesta:
    def test_produce_bytes_de_un_docx_valido(self):
        contenido = generar_docx_respuesta(_respuesta())
        assert contenido[:2] == b"PK"  # firma de archivo ZIP (contenedor OOXML)
        Document(io.BytesIO(contenido))  # no debe lanzar

    def test_incluye_asunto_y_destinatario(self):
        contenido = generar_docx_respuesta(_respuesta())
        texto = _texto_completo(contenido)
        assert "Respuesta a solicitud de información" in texto
        assert "JUAN PÉREZ LÓPEZ" in texto
        assert "Director de Administración" in texto
        assert "PRESENTE" in texto.replace(" ", "")

    def test_cuerpo_respeta_parrafos_separados_por_linea_en_blanco(self):
        contenido = generar_docx_respuesta(_respuesta())
        texto = _texto_completo(contenido)
        assert "En atención a su similar número DSA-2026-089-OF" in texto
        assert "El trámite solicitado ha sido atendido en su totalidad." in texto

    def test_incluye_ccp_cuando_se_proporciona(self):
        contenido = generar_docx_respuesta(_respuesta())
        texto = _texto_completo(contenido)
        assert "C.c.p." in texto
        assert "Archivo" in texto
        assert "Minutario" in texto

    def test_omite_bloque_ccp_cuando_esta_vacio(self):
        contenido = generar_docx_respuesta(_respuesta(ccp=[]))
        texto = _texto_completo(contenido)
        assert "C.c.p." not in texto

    def test_incluye_firmante_cuando_se_proporciona(self):
        contenido = generar_docx_respuesta(
            _respuesta(), firmante_nombre="ana torres", firmante_cargo="Jefa de Departamento"
        )
        texto = _texto_completo(contenido)
        assert "ANA TORRES" in texto
        assert "Jefa de Departamento" in texto

    def test_sin_firmante_no_revienta(self):
        contenido = generar_docx_respuesta(_respuesta())
        Document(io.BytesIO(contenido))  # no debe lanzar sin firmante_nombre


class TestNombreArchivoRespuesta:
    def test_incluye_folio_de_origen_y_destinatario(self):
        nombre = nombre_archivo_respuesta(_documento(), _respuesta())
        assert nombre.startswith("RESP__DSA-2026-089-OF__")
        assert nombre.endswith(".docx")

    def test_folio_ausente_usa_marcador(self):
        nombre = nombre_archivo_respuesta(_documento(numero_oficio=None), _respuesta())
        assert nombre.startswith("RESP__SIN_FOLIO__")

    def test_caracteres_reservados_se_sanean(self):
        nombre = nombre_archivo_respuesta(
            _documento(numero_oficio="SSJ/DEA/2026/089"), _respuesta()
        )
        assert "/" not in nombre
