"""
SISTEMA OFICIALIA-DIGITAL-DSA (reconstrucción 100% Python)
==========================================================
core/doc_generator.py — Generación del .docx de la respuesta a un oficio.

Convierte una `RespuestaOficio` ya APROBADA por HITL (ver
core/pipeline.py::FlujoDocumental.generar_documento_respuesta) en un
documento Word editable, listo para imprimirse, firmarse o turnarse por
RPA. Deliberadamente NO produce PDF ni firma nada: el funcionario conserva
el archivo editable para ajustes de último momento antes de imprimir.

Plantilla con membrete oficial: si `plantilla_path` apunta a un .docx con
los logotipos/pie de página institucionales, el contenido se agrega al
final de esa plantilla en vez de partir de un documento en blanco.
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from core.models import DocumentoRegistro, RespuestaOficio

#: Caracteres prohibidos en nombres de archivo (mismo criterio que core.models).
_CARACTERES_RESERVADOS_RE = re.compile(r"[/\\:*?\"<>|]")


def nombre_archivo_respuesta(documento: DocumentoRegistro, respuesta: RespuestaOficio) -> str:
    """
    Nombre sugerido para el .docx de contestación:
        RESP__[FOLIO_ORIGEN]__[DESTINATARIO].docx

    :param documento: el oficio de origen que se está contestando.
    :param respuesta: contenido aprobado de la contestación.
    """
    folio_origen = documento.numero_oficio or "SIN_FOLIO"
    folio_origen = _CARACTERES_RESERVADOS_RE.sub("-", folio_origen.strip()) or "SIN_FOLIO"
    destinatario = _CARACTERES_RESERVADOS_RE.sub(
        "-", respuesta.destinatario_nombre[:30].strip()
    ).replace(" ", "_")
    destinatario = destinatario or "SIN_DESTINATARIO"
    return f"RESP__{folio_origen}__{destinatario}.docx"


def generar_docx_respuesta(
    respuesta: RespuestaOficio,
    *,
    firmante_nombre: str = "",
    firmante_cargo: str = "",
    plantilla_path: Optional[Path] = None,
) -> bytes:
    """
    Construye el documento Word de la contestación y lo devuelve como bytes
    (listo para persistir en storage/ o servir por descarga).

    :param firmante_nombre: nombre de quien firma (capturado en
        `PeticionRespuesta.firmante_nombre`, no forma parte del contrato
        `RespuestaOficio` porque la IA no debe inventar quién firma).
    :param firmante_cargo: cargo de quien firma.
    :param plantilla_path: .docx base con membrete institucional editable
        (logotipos, pie de página). Si se omite o el archivo no existe, se
        genera un documento en blanco con `python-docx`.
    """
    if plantilla_path is not None and Path(plantilla_path).is_file():
        doc = Document(str(plantilla_path))
    else:
        doc = Document()

    # ---- Asunto ----
    p_asunto = doc.add_paragraph()
    p_asunto.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_asunto.add_run(f"ASUNTO: {respuesta.asunto}").bold = True

    doc.add_paragraph()  # separación visual

    # ---- Destinatario ----
    p_dest = doc.add_paragraph()
    p_dest.add_run(respuesta.destinatario_nombre).bold = True
    if respuesta.destinatario_cargo:
        p_dest.add_run("\n" + respuesta.destinatario_cargo)
    if respuesta.destinatario_dependencia:
        p_dest.add_run("\n" + respuesta.destinatario_dependencia)
    p_dest.add_run("\nP R E S E N T E.").bold = True

    doc.add_paragraph()

    # ---- Cuerpo (párrafos separados por línea en blanco) ----
    for parrafo in respuesta.cuerpo_respuesta.split("\n\n"):
        parrafo = parrafo.strip()
        if not parrafo:
            continue
        p_cuerpo = doc.add_paragraph(parrafo)
        p_cuerpo.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # ---- Despedida y firma ----
    doc.add_paragraph()
    p_despedida = doc.add_paragraph(respuesta.despedida)
    p_despedida.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    doc.add_paragraph()
    doc.add_paragraph()
    p_atentamente = doc.add_paragraph("A T E N T A M E N T E")
    p_atentamente.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if firmante_nombre.strip():
        # Espacio en blanco para la rúbrica manuscrita antes del nombre impreso.
        for _ in range(3):
            doc.add_paragraph()
        p_firma = doc.add_paragraph(firmante_nombre.strip().upper())
        p_firma.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_firma.runs[0].bold = True
        if firmante_cargo.strip():
            p_cargo = doc.add_paragraph(firmante_cargo.strip())
            p_cargo.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ---- C.c.p. ----
    if respuesta.ccp:
        p_ccp = doc.add_paragraph()
        p_ccp.paragraph_format.space_before = Pt(20)
        p_ccp.add_run("C.c.p. -").font.size = Pt(8)
        for area in respuesta.ccp:
            p_ccp.add_run(f"\n• {area}").font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
